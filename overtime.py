"""
Overtime (ზეგანაკვეთური სამუშაოები) — დამოუკიდებელი მოდული.

იზოლაცია:
  * მონაცემები ინახება ცალკე Postgres სქემაში `overtime` (იმავე Neon ბაზაში),
    ან სულ სხვა ბაზაში, თუ დაყენებულია OVERTIME_DATABASE_URL.
  * ძირითადი მოდულის ცხრილებს (records, settings, activity_log) მოდული არ ცვლის —
    `records`-ს მხოლოდ კითხულობს ავტომატური შევსებისთვის (ID / განწესი).
  * მოდულის ჩატვირთვის ან მისი ბაზის შეცდომა ძირითად აპლიკაციას არ აჩერებს.

დღე / ღამე: დღე = 06:00–22:00, ღამე = 22:00–06:00 (overtime.xlsx-ის M/N ფორმულების იდენტური ლოგიკა,
შემოწმებულია ფაილის 4 826 სტრიქონზე). შუაღამის გადაკვეთისას დასრულება მომდევნო დღეს ეკუთვნის.
"""
import io
import os
import re
import time
import json
import uuid
from datetime import datetime, date as date_cls, time as time_cls, timedelta, timezone
from contextlib import contextmanager

from flask import Blueprint, request, jsonify, send_file, render_template

import psycopg2
from psycopg2.extras import RealDictCursor

bp = Blueprint("overtime", __name__)

SCHEMA = "overtime"
DAY_START, DAY_END = 6 * 60, 22 * 60          # 06:00 – 22:00
WEEKDAYS = ["ორშაბათი", "სამშაბათი", "ოთხშაბათი", "ხუთშაბათი", "პარასკევი", "შაბათი", "კვირა"]
ROLES = ("none", "operator", "manager")
COMMENT_CHOICES = ["შაბათი/კვირა", "უქმე", "იმუშავა მეორე დღეს"]
GANWESI_RE = re.compile(r"^\d{1,3}/\d{1,6}$")
PID_RE = re.compile(r"^\d{11}$")
MAX_BATCH = 60            # თანამშრომელი ერთ შენახვაზე
MAX_IMPORT_BATCH = 500

_deps = None              # dict: current_user, client_ip, db_cursor (main), load_config
_ready = {"ok": False, "error": ""}


# ---------------------------------------------------------------------------
# DB
# ---------------------------------------------------------------------------
def _ot_url():
    return os.environ.get("OVERTIME_DATABASE_URL") or os.environ.get("POSTGRES_URL")


@contextmanager
def ot_cursor(commit=False):
    url = _ot_url()
    if not url:
        raise RuntimeError("database url is not configured")
    conn = psycopg2.connect(url, cursor_factory=RealDictCursor, connect_timeout=8)
    try:
        cur = conn.cursor()
        yield cur
        if commit:
            conn.commit()
        cur.close()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


T_ENTRIES = f"{SCHEMA}.entries"
T_AUDIT = f"{SCHEMA}.audit"


def init_schema():
    """არადესტრუქციული: მხოლოდ CREATE ... IF NOT EXISTS. არსებულ ცხრილებს არ ეხება."""
    with ot_cursor(commit=True) as cur:
        cur.execute(f"CREATE SCHEMA IF NOT EXISTS {SCHEMA};")
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {T_ENTRIES} (
                id BIGSERIAL PRIMARY KEY,
                uid VARCHAR(40) NOT NULL UNIQUE,
                work_date DATE NOT NULL,
                emp_key VARCHAR(120) NOT NULL,
                emp_name VARCHAR(120) NOT NULL,
                emp_pid VARCHAR(20) NOT NULL DEFAULT '',
                emp_position VARCHAR(80) NOT NULL DEFAULT '',
                brigade VARCHAR(40) NOT NULL DEFAULT '',
                work_id VARCHAR(40) NOT NULL DEFAULT '',
                ganwesi VARCHAR(40) NOT NULL DEFAULT '',
                object_name VARCHAR(300) NOT NULL DEFAULT '',
                address VARCHAR(400) NOT NULL DEFAULT '',
                start_min INTEGER NOT NULL,
                end_min INTEGER NOT NULL,
                day_min INTEGER NOT NULL,
                night_min INTEGER NOT NULL,
                comment VARCHAR(300) NOT NULL DEFAULT '',
                note VARCHAR(500) NOT NULL DEFAULT '',
                source VARCHAR(20) NOT NULL DEFAULT 'manual',
                source_ref VARCHAR(100) NOT NULL DEFAULT '',
                autofilled VARCHAR(300) NOT NULL DEFAULT '',
                created_by VARCHAR(100) NOT NULL DEFAULT '',
                created_ts TIMESTAMPTZ NOT NULL,
                updated_by VARCHAR(100) NOT NULL DEFAULT '',
                updated_ts TIMESTAMPTZ,
                deleted_ts TIMESTAMPTZ,
                deleted_by VARCHAR(100) NOT NULL DEFAULT ''
            );
        """)
        # ზუსტი დუბლიკატი (იგივე ადამიანი, დღე, დაწყება, დასრულება) ბაზის დონეზეც შეუძლებელია
        cur.execute(f"""CREATE UNIQUE INDEX IF NOT EXISTS ot_entries_exact ON {T_ENTRIES}
                        (emp_key, work_date, start_min, end_min) WHERE deleted_ts IS NULL;""")
        cur.execute(f"CREATE INDEX IF NOT EXISTS ot_entries_date ON {T_ENTRIES} (work_date);")
        cur.execute(f"CREATE INDEX IF NOT EXISTS ot_entries_emp ON {T_ENTRIES} (emp_key, work_date);")
        cur.execute(f"CREATE INDEX IF NOT EXISTS ot_entries_ganwesi ON {T_ENTRIES} (ganwesi);")
        cur.execute(f"CREATE INDEX IF NOT EXISTS ot_entries_workid ON {T_ENTRIES} (work_id);")
        cur.execute(f"""
            CREATE TABLE IF NOT EXISTS {T_AUDIT} (
                id BIGSERIAL PRIMARY KEY,
                ts TIMESTAMPTZ NOT NULL,
                username VARCHAR(100) NOT NULL DEFAULT '',
                ip VARCHAR(64) NOT NULL DEFAULT '',
                action VARCHAR(40) NOT NULL,
                entry_uid VARCHAR(40) NOT NULL DEFAULT '',
                details TEXT NOT NULL DEFAULT '',
                reason VARCHAR(300) NOT NULL DEFAULT ''
            );
        """)
        cur.execute(f"CREATE INDEX IF NOT EXISTS ot_audit_ts ON {T_AUDIT} (ts);")
    # აუდიტი — მხოლოდ დამატება (Postgres). SQLite-ის სატესტო გარემოში უბრალოდ გამოტოვდება.
    try:
        with ot_cursor(commit=True) as cur:
            cur.execute(f"""
                CREATE OR REPLACE FUNCTION {SCHEMA}.audit_append_only() RETURNS trigger AS $$
                BEGIN RAISE EXCEPTION 'overtime.audit is append-only'; END; $$ LANGUAGE plpgsql;
            """)
            cur.execute(f"DROP TRIGGER IF EXISTS audit_no_change ON {T_AUDIT};")
            cur.execute(f"""CREATE TRIGGER audit_no_change BEFORE UPDATE OR DELETE ON {T_AUDIT}
                            FOR EACH ROW EXECUTE FUNCTION {SCHEMA}.audit_append_only();""")
            cur.execute(f"DROP TRIGGER IF EXISTS audit_no_truncate ON {T_AUDIT};")
            cur.execute(f"""CREATE TRIGGER audit_no_truncate BEFORE TRUNCATE ON {T_AUDIT}
                            FOR EACH STATEMENT EXECUTE FUNCTION {SCHEMA}.audit_append_only();""")
    except Exception as e:
        print("overtime audit trigger skipped:", type(e).__name__)


RETRY_AFTER = 30   # წამი: მიუწვდომელ ბაზას ყოველ მოთხოვნაზე არ ველოდებით


def ensure_ready():
    """მოდულის ბაზა მზადაა? პირველი Overtime-მოთხოვნისას ქმნის სქემას (IF NOT EXISTS).
    შეცდომა მხოლოდ Overtime-ს ეხება; ხელახლა ცდა — არაუმეტეს 30 წამში ერთხელ."""
    if _ready["ok"]:
        return True
    if _ready.get("failed_at") and time.monotonic() - _ready["failed_at"] < RETRY_AFTER:
        return False
    try:
        init_schema()
        _ready.update(ok=True, error="", failed_at=0)
    except Exception as e:
        _ready.update(error=type(e).__name__, failed_at=time.monotonic())
        _log_error("init", e)
    return _ready["ok"]


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def _mask(text):
    """ლოგში პირადი ნომრები და გრძელი ციფრები არ უნდა მოხვდეს."""
    return re.sub(r"\d{6,}", "***", str(text))


def _log_error(where, e):
    print(f"OVERTIME ERROR [{where}] {type(e).__name__}: {_mask(e)[:300]}")


def _err(msg, status=400, **extra):
    body = {"error": msg}
    body.update(extra)
    return jsonify(body), status


def _now():
    return datetime.now(timezone.utc)


def _user():
    return _deps["current_user"]() if _deps else None


def role_of(user):
    r = str((user or {}).get("ot_role") or "none")
    return r if r in ROLES else "none"


def _require(level):
    """level: 'operator' (ოპერატორი ან მენეჯერი) ან 'manager'. უფლება მოწმდება სერვერზე."""
    u = _user()
    if not u:
        return None, _err("საჭიროა სისტემაში შესვლა", 401, login_required=True)
    r = role_of(u)
    if r == "none" or (level == "manager" and r != "manager"):
        return None, _err("ამ მოქმედებისთვის Overtime-ის უფლება არ გაქვთ", 403)
    if not ensure_ready():
        return None, _err("Overtime-ის ბაზა დროებით მიუწვდომელია. ძირითადი მოდული ჩვეულებრივ მუშაობს.", 503)
    return u, None


def norm_ganwesi(v):
    return re.sub(r"\s+", "", str(v or "")).strip()[:40]


def norm_work_id(v):
    s = str(v or "").strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s[:40]


def norm_name(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def emp_key(name, pid):
    pid = re.sub(r"\D", "", str(pid or ""))
    return f"pid:{pid}" if pid else "name:" + norm_name(name).lower()


def parse_hhmm(v):
    m = re.fullmatch(r"(\d{1,2}):(\d{2})", str(v or "").strip())
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    if h > 23 or mi > 59:
        return None
    return h * 60 + mi


def fmt_hhmm(m):
    return f"{m // 60:02d}:{m % 60:02d}"


def split_day_night(start_min, end_min):
    """(დღის წუთები, ღამის წუთები). შუაღამის გადაკვეთა — დასრულება მომდევნო დღეზე."""
    s, e = start_min, end_min
    if e <= s:
        e += 1440
    day = 0
    for a, b in ((DAY_START, DAY_END), (DAY_START + 1440, DAY_END + 1440)):
        day += max(0, min(e, b) - max(s, a))
    return day, (e - s) - day


def parse_date(v):
    s = str(v or "").strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        return None


def _d(v):
    """DATE მნიშვნელობა ორივე ბაზიდან (Postgres: date, SQLite: str) → date."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date_cls):
        return v
    return parse_date(v)


def _ts(v):
    if isinstance(v, datetime):
        return (v if v.tzinfo else v.replace(tzinfo=timezone.utc)).isoformat()
    return str(v or "")


def _abs_interval(d, s, e):
    base = d.toordinal() * 1440
    end = e if e > s else e + 1440
    return base + s, base + end


def entry_public(r):
    d = _d(r["work_date"])
    return {
        "id": r["id"], "uid": r["uid"], "work_date": d.isoformat() if d else "",
        "weekday": WEEKDAYS[d.weekday()] if d else "",
        "emp_name": r["emp_name"], "emp_pid": r["emp_pid"], "emp_position": r["emp_position"],
        "brigade": r["brigade"], "work_id": r["work_id"], "ganwesi": r["ganwesi"],
        "object_name": r["object_name"], "address": r["address"],
        "start": fmt_hhmm(r["start_min"]), "end": fmt_hhmm(r["end_min"]),
        "day_min": r["day_min"], "night_min": r["night_min"],
        "crosses_midnight": r["end_min"] <= r["start_min"],
        "comment": r["comment"], "note": r["note"], "source": r["source"],
        "autofilled": [x for x in str(r.get("autofilled") or "").split(",") if x],
        "created_by": r["created_by"], "created_ts": _ts(r["created_ts"]),
        "updated_by": r["updated_by"], "updated_ts": _ts(r.get("updated_ts")),
        "deleted": bool(r.get("deleted_ts")), "deleted_by": r.get("deleted_by") or "",
    }


def audit(cur, action, entry_uid="", details="", reason="", username=None):
    u = _user()
    cur.execute(f"""INSERT INTO {T_AUDIT} (ts, username, ip, action, entry_uid, details, reason)
                    VALUES (%s, %s, %s, %s, %s, %s, %s);""",
                (_now(), str(username if username is not None else (u or {}).get("username", ""))[:100],
                 _deps["client_ip"]() if _deps else "", action[:40], str(entry_uid)[:40],
                 str(details)[:4000], str(reason or "")[:300]))


# ---------------------------------------------------------------------------
# ვალიდაცია
# ---------------------------------------------------------------------------
EDITABLE_OPERATOR = {"start", "end", "comment", "note"}
EDITABLE_MANAGER = {"work_date", "brigade", "work_id", "ganwesi", "object_name", "address",
                    "start", "end", "comment", "note", "emp_name", "emp_pid", "emp_position"}


def _known_brigade_of(pid):
    """პირადი ნომრით — რომელ ბრიგადაშია ძირითად პარამეტრებში (შესაბამისობის შესამოწმებლად)."""
    if not pid or not _deps:
        return None
    try:
        cfg = _deps["load_config"]()
    except Exception:
        return None
    for b, data in (cfg.get("brigades") or {}).items():
        for m in data.get("members", []):
            if str(m.get("personal_id") or "") == pid:
                return str(b)
    return None


def validate_common(c):
    """საერთო ველები. აბრუნებს (cleaned, errors, warnings)."""
    errors, warnings = [], []
    out = {}
    d = parse_date(c.get("work_date"))
    if not d:
        errors.append("მიუთითეთ სამუშაოს თარიღი")
    else:
        today = datetime.now(timezone(timedelta(hours=4))).date()
        if d < date_cls(2020, 1, 1) or d > today + timedelta(days=1):
            errors.append("თარიღი დასაშვებ შუალედში არ არის")
    out["work_date"] = d
    s, e = parse_hhmm(c.get("start")), parse_hhmm(c.get("end"))
    if s is None or e is None:
        errors.append("მიუთითეთ დაწყებისა და დასრულების დრო (სს:წწ)")
    elif s == e:
        errors.append("დაწყება და დასრულება ერთნაირია — ხანგრძლივობა 0-ია")
    else:
        dur = (e - s) % 1440
        if dur > 16 * 60:
            warnings.append(f"ხანგრძლივობა {dur // 60}სთ {dur % 60:02d}წ — 16 საათზე მეტია")
    out["start_min"], out["end_min"] = s, e
    out["work_id"] = norm_work_id(c.get("work_id"))
    out["ganwesi"] = norm_ganwesi(c.get("ganwesi"))
    if not out["work_id"] and not out["ganwesi"]:
        errors.append("მიუთითეთ სამუშაოს ID ან განწესის ნომერი")
    if out["ganwesi"] and not GANWESI_RE.match(out["ganwesi"]):
        warnings.append(f"განწესის ნომერი „{out['ganwesi']}“ არასტანდარტული ფორმატისაა (მოსალოდნელია მაგ. 12/0415)")
    if out["work_id"] and not out["work_id"].isdigit():
        warnings.append(f"სამუშაოს ID „{out['work_id']}“ რიცხვი არ არის")
    for k, lim in (("brigade", 40), ("object_name", 300), ("address", 400), ("comment", 300), ("note", 500)):
        v = str(c.get(k) or "").strip()
        if len(v) > lim:
            errors.append(f"ველი „{k}“ ძალიან გრძელია")
        out[k] = v[:lim]
    out["source"] = str(c.get("source") or "manual")[:20]
    if out["source"] not in ("manual", "reports", "overtime", "import", "restore"):
        out["source"] = "manual"
    out["source_ref"] = str(c.get("source_ref") or "")[:100]
    af = c.get("autofilled") or []
    out["autofilled"] = ",".join(str(x)[:30] for x in af if isinstance(x, str))[:300] if isinstance(af, list) else ""
    return out, errors, warnings


def validate_employee(e, brigade):
    errors, warnings = [], []
    name = norm_name(e.get("name"))
    pid = re.sub(r"\s", "", str(e.get("pid") or ""))
    pos = str(e.get("position") or "").strip()[:80]
    if not name:
        errors.append("თანამშრომლის სახელი და გვარი სავალდებულოა")
    elif len(name) > 120:
        errors.append("სახელი ძალიან გრძელია")
    if pid and not PID_RE.match(pid):
        errors.append(f"{name}: პირადი ნომერი 11 ციფრი უნდა იყოს")
    if not pid and name:
        warnings.append(f"{name}: პირადი ნომერი არ არის მითითებული — დუბლიკატის კონტროლი სახელით მოხდება")
    known = _known_brigade_of(pid) if pid else None
    if known and brigade and str(brigade).strip() != known:
        warnings.append(f"{name} ძირითად სიაში ბრიგადა {known}-შია (ჩანაწერში: {brigade})")
    return {"name": name, "pid": pid, "position": pos, "key": emp_key(name, pid)}, errors, warnings


def find_conflicts(cur, key, d, s, e, ganwesi, work_id, exclude_id=None):
    """ზუსტი დუბლიკატი (ბლოკავს), დროის გადაფარვა და იმავე სამუშაოს ჩანაწერი (დადასტურებით)."""
    cur.execute(f"""SELECT * FROM {T_ENTRIES} WHERE emp_key = %s AND work_date >= %s AND work_date <= %s
                    AND deleted_ts IS NULL;""",
                (key, d - timedelta(days=1), d + timedelta(days=1)))
    exact, overlap, same_work = [], [], []
    a0, a1 = _abs_interval(d, s, e)
    for r in cur.fetchall():
        if exclude_id is not None and int(r["id"]) == int(exclude_id):
            continue
        rd = _d(r["work_date"])
        if rd == d and r["start_min"] == s and r["end_min"] == e:
            exact.append(entry_public(r))
            continue
        b0, b1 = _abs_interval(rd, r["start_min"], r["end_min"])
        if a0 < b1 and b0 < a1:
            ov = min(a1, b1) - max(a0, b0)
            overlap.append(dict(entry_public(r), overlap_min=ov))
        elif rd == d and ((ganwesi and r["ganwesi"] == ganwesi) or (work_id and r["work_id"] == work_id)):
            same_work.append(entry_public(r))
    return exact, overlap, same_work


# ---------------------------------------------------------------------------
# Routes: page + session info
# ---------------------------------------------------------------------------
@bp.route("/overtime")
def page():
    return render_template("overtime.html")


@bp.route("/api/ot/status")
def status():
    u = _user()
    if not u:
        return _err("საჭიროა სისტემაში შესვლა", 401, login_required=True)
    r = role_of(u)
    return jsonify({"role": r, "ready": ensure_ready() if r != "none" else False,
                    "comment_choices": COMMENT_CHOICES,
                    "can": {"create": r in ("operator", "manager"), "edit_all": r == "manager",
                            "delete": r == "manager", "export": r == "manager", "import": r == "manager"}})


# ---------------------------------------------------------------------------
# ავტომატური შევსება: ID / განწესი → ძირითადი მოდულის ჩანაწერები (მხოლოდ წაკითხვა)
# ---------------------------------------------------------------------------
def _candidate_from_record(row, via):
    rec = row["data"] if isinstance(row.get("data"), dict) else {}
    members = []
    for m in rec.get("members") or []:
        if not isinstance(m, dict):
            continue
        absent = bool(m.get("absent") or str(m.get("note") or "").strip())
        members.append({"name": norm_name(m.get("name")), "pid": str(m.get("personal_id") or ""),
                        "position": str(m.get("position") or ""), "absent": absent,
                        "absent_reason": str(m.get("note") or "").strip()})
    times = [f"{w.get('start')}–{w.get('end')}" for w in rec.get("works") or [] if isinstance(w, dict) and w.get("start") and w.get("end")]
    return {
        "via": via, "source": "reports", "ref": str(row.get("db_id") or ""),
        "work_id": norm_work_id(rec.get("id")), "ganwesi": norm_ganwesi(rec.get("ganwesi")),
        "work_date": str(rec.get("date") or "")[:10], "brigade": str(rec.get("brigade") or ""),
        "city": str(rec.get("city") or ""), "object_name": str(rec.get("object_name") or ""),
        "address": str(rec.get("address") or ""), "members": members, "work_times": times,
    }


def _lookup_reports(work_id, ganwesi):
    by_id, by_g = [], []
    with _deps["db_cursor"]() as cur:
        if work_id:
            cur.execute("SELECT db_id, data FROM records WHERE id = %s ORDER BY created_at DESC, db_id DESC LIMIT 30;", (work_id,))
            by_id = cur.fetchall()
        if ganwesi:
            cur.execute("SELECT db_id, data FROM records WHERE REPLACE(data->>'ganwesi', ' ', '') = %s "
                        "ORDER BY created_at DESC, db_id DESC LIMIT 30;", (ganwesi,))
            by_g = cur.fetchall()
    return by_id, by_g


def _lookup_history(cur, work_id, ganwesi):
    conds, params = [], []
    if work_id:
        conds.append("work_id = %s"); params.append(work_id)
    if ganwesi:
        conds.append("ganwesi = %s"); params.append(ganwesi)
    if not conds:
        return []
    cur.execute(f"SELECT * FROM {T_ENTRIES} WHERE ({' OR '.join(conds)}) AND deleted_ts IS NULL "
                f"ORDER BY work_date DESC, id DESC LIMIT 200;", tuple(params))
    groups = {}
    for r in cur.fetchall():
        d = _d(r["work_date"])
        k = (r["work_id"], r["ganwesi"], d, r["brigade"], r["object_name"], r["address"])
        g = groups.setdefault(k, {"via": "history", "source": "overtime", "ref": "", "work_id": r["work_id"],
                                  "ganwesi": r["ganwesi"], "work_date": d.isoformat() if d else "",
                                  "brigade": r["brigade"], "city": "", "object_name": r["object_name"],
                                  "address": r["address"], "members": [], "work_times": []})
        if not any(m["pid"] == r["emp_pid"] and m["name"] == r["emp_name"] for m in g["members"]):
            g["members"].append({"name": r["emp_name"], "pid": r["emp_pid"], "position": r["emp_position"],
                                 "absent": False, "absent_reason": ""})
    return list(groups.values())[:10]


@bp.route("/api/ot/lookup")
def lookup():
    u, err = _require("operator")
    if err:
        return err
    work_id = norm_work_id(request.args.get("work_id"))
    ganwesi = norm_ganwesi(request.args.get("ganwesi"))
    if not work_id and not ganwesi:
        return _err("შეიყვანეთ სამუშაოს ID ან განწესის ნომერი")
    try:
        by_id, by_g = _lookup_reports(work_id, ganwesi)
    except Exception as e:
        _log_error("lookup.reports", e)
        return _err("ძირითად ჩანაწერებში ძიება ვერ მოხერხდა. შეავსეთ ხელით ან სცადეთ მოგვიანებით.", 503,
                    status="unavailable")
    cands, status_ = [], "none"
    if work_id and ganwesi:
        both = [r for r in by_id if norm_ganwesi((r["data"] or {}).get("ganwesi")) == ganwesi]
        if both:
            status_, cands = "match", [_candidate_from_record(r, "both") for r in both]
        elif by_id or by_g:
            status_ = "mismatch"
            cands = [_candidate_from_record(r, "id") for r in by_id] + [_candidate_from_record(r, "ganwesi") for r in by_g]
    elif by_id or by_g:
        status_ = "found"
        cands = [_candidate_from_record(r, "id" if work_id else "ganwesi") for r in (by_id or by_g)]
    history = []
    try:
        with ot_cursor() as cur:
            history = _lookup_history(cur, work_id, ganwesi)
    except Exception as e:
        _log_error("lookup.history", e)
    if status_ == "none" and history:
        status_ = "history"
    messages = {
        "match": "ID და განწესი ერთმანეთს ემთხვევა — მონაცემები ჩანაწერიდან შეივსება.",
        "found": "ჩანაწერი მოიძებნა.",
        "mismatch": "ID და განწესი სხვადასხვა ჩანაწერს ეკუთვნის. გადაამოწმეთ და აირჩიეთ სწორი, ან შეავსეთ ხელით.",
        "history": "ძირითად ჩანაწერებში ვერ მოიძებნა, მაგრამ არის Overtime-ის წინა ჩანაწერები.",
        "none": ("მითითებული ID ვერ მოიძებნა. გთხოვთ, ხელით შეავსოთ სამუშაოს საჭირო ინფორმაცია." if work_id and not ganwesi else
                 "მითითებული განწესი ვერ მოიძებნა. გთხოვთ, ხელით შეავსოთ სამუშაოს საჭირო ინფორმაცია." if ganwesi and not work_id else
                 "მითითებული ID და განწესი ვერ მოიძებნა. გთხოვთ, ხელით შეავსოთ სამუშაოს საჭირო ინფორმაცია."),
    }
    return jsonify({"status": status_, "message": messages[status_], "query": {"work_id": work_id, "ganwesi": ganwesi},
                    "candidates": cands[:30], "history": history})


BRIGADE_CITIES = ("თბილისი",)   # ამ ეტაპზე ბრიგადის შემადგენლობა მხოლოდ თბილისის ბრიგადებისთვის


@bp.route("/api/ot/brigades")
def brigades():
    """ხელით შევსებისთვის: ბრიგადის ნომერი → შემადგენლობა (ძირითადი პარამეტრებიდან, მხოლოდ წაკითხვა)."""
    u, err = _require("operator")
    if err:
        return err
    out = []
    try:
        cfg = _deps["load_config"]()
        for b, data in (cfg.get("brigades") or {}).items():
            members = [{"name": norm_name(m.get("name")), "pid": str(m.get("personal_id") or ""),
                        "position": str(m.get("position") or "")}
                       for m in data.get("members", []) if m.get("city") in BRIGADE_CITIES and m.get("name")]
            if members:
                out.append({"brigade": str(b), "members": members})
    except Exception as e:
        _log_error("brigades", e)
        return _err("ბრიგადების სია ვერ ჩაიტვირთა", 500)
    out.sort(key=lambda x: (0, int(x["brigade"])) if x["brigade"].isdigit() else (1, x["brigade"]))
    return jsonify({"cities": list(BRIGADE_CITIES), "brigades": out})


@bp.route("/api/ot/employees")
def employees():
    """თანამშრომლების ძიება: ძირითადი მოდულის შემადგენლობა + Overtime-ის წინა ჩანაწერები."""
    u, err = _require("operator")
    if err:
        return err
    q = norm_name(request.args.get("q")).lower()
    if len(q) < 2:
        return jsonify({"rows": []})
    out, seen = [], set()
    try:
        cfg = _deps["load_config"]()
        for b, data in (cfg.get("brigades") or {}).items():
            for m in data.get("members", []):
                name, pid = norm_name(m.get("name")), str(m.get("personal_id") or "")
                if q in name.lower() or (q.isdigit() and q in pid):
                    k = emp_key(name, pid)
                    if k not in seen:
                        seen.add(k)
                        out.append({"name": name, "pid": pid, "position": m.get("position", ""), "brigade": str(b), "source": "ბრიგადა"})
    except Exception as e:
        _log_error("employees.config", e)
    try:
        with ot_cursor() as cur:
            like = f"%{q}%"
            cur.execute(f"""SELECT emp_name, emp_pid, emp_position, MAX(brigade) AS brigade FROM {T_ENTRIES}
                            WHERE (LOWER(emp_name) LIKE %s OR emp_pid LIKE %s) AND deleted_ts IS NULL
                            GROUP BY emp_name, emp_pid, emp_position LIMIT 40;""", (like, like))
            for r in cur.fetchall():
                k = emp_key(r["emp_name"], r["emp_pid"])
                if k not in seen:
                    seen.add(k)
                    out.append({"name": r["emp_name"], "pid": r["emp_pid"], "position": r["emp_position"],
                                "brigade": r["brigade"] or "", "source": "Overtime"})
    except Exception as e:
        _log_error("employees.history", e)
    return jsonify({"rows": out[:25]})


# ---------------------------------------------------------------------------
# ჩანაწერები
# ---------------------------------------------------------------------------
def _filters():
    a = request.args
    where, params = [], []
    if a.get("deleted") == "1":
        where.append("deleted_ts IS NOT NULL")
    else:
        where.append("deleted_ts IS NULL")
    for key, col in (("from", ">="), ("to", "<=")):
        d = parse_date(a.get(key))
        if d:
            where.append(f"work_date {col} %s"); params.append(d)
    if a.get("emp"):
        q = norm_name(a["emp"]).lower()
        where.append("(LOWER(emp_name) LIKE %s OR emp_pid LIKE %s)"); params += [f"%{q}%", f"%{q}%"]
    if a.get("brigade"):
        where.append("brigade = %s"); params.append(str(a["brigade"]).strip())
    if a.get("work_id"):
        where.append("work_id = %s"); params.append(norm_work_id(a["work_id"]))
    if a.get("ganwesi"):
        where.append("ganwesi = %s"); params.append(norm_ganwesi(a["ganwesi"]))
    if a.get("q"):
        q = f"%{str(a['q']).strip().lower()}%"
        where.append("(LOWER(object_name) LIKE %s OR LOWER(address) LIKE %s OR LOWER(comment) LIKE %s OR LOWER(note) LIKE %s)")
        params += [q, q, q, q]
    return " WHERE " + " AND ".join(where), params


@bp.route("/api/ot/entries", methods=["GET"])
def list_entries():
    u, err = _require("operator")
    if err:
        return err
    sql_where, params = _filters()
    try:
        limit = max(1, min(int(request.args.get("limit", 100)), 500))
        offset = max(0, int(request.args.get("offset", 0)))
    except ValueError:
        limit, offset = 100, 0
    try:
        with ot_cursor() as cur:
            cur.execute(f"""SELECT COUNT(*) AS n, COALESCE(SUM(day_min), 0) AS d, COALESCE(SUM(night_min), 0) AS nt,
                            COUNT(DISTINCT emp_key) AS people FROM {T_ENTRIES}{sql_where};""", tuple(params))
            agg = cur.fetchone()
            cur.execute(f"SELECT * FROM {T_ENTRIES}{sql_where} ORDER BY work_date DESC, ganwesi, emp_name, start_min "
                        f"LIMIT %s OFFSET %s;", tuple(params) + (limit, offset))
            rows = [entry_public(r) for r in cur.fetchall()]
            cur.execute(f"SELECT DISTINCT brigade FROM {T_ENTRIES} WHERE deleted_ts IS NULL AND brigade <> '' ORDER BY brigade;")
            brigades = [r["brigade"] for r in cur.fetchall()]
    except Exception as e:
        _log_error("list", e)
        return _err("ჩანაწერების ჩატვირთვა ვერ მოხერხდა", 500)
    return jsonify({"rows": rows, "total": int(agg["n"]), "day_min": int(agg["d"]), "night_min": int(agg["nt"]),
                    "people": int(agg["people"]), "brigades": brigades, "limit": limit, "offset": offset})


@bp.route("/api/ot/check")
def check_work():
    """ფორმისთვის: ამ სამუშაოზე (განწესი/ID + თარიღი) უკვე შეტანილი ჩანაწერები."""
    u, err = _require("operator")
    if err:
        return err
    g, wid, d = norm_ganwesi(request.args.get("ganwesi")), norm_work_id(request.args.get("work_id")), parse_date(request.args.get("date"))
    conds, params = [], []
    if g:
        conds.append("ganwesi = %s"); params.append(g)
    if wid:
        conds.append("work_id = %s"); params.append(wid)
    if not conds:
        return jsonify({"rows": []})
    sql = f"SELECT * FROM {T_ENTRIES} WHERE ({' OR '.join(conds)}) AND deleted_ts IS NULL"
    if d:
        sql += " AND work_date = %s"; params.append(d)
    try:
        with ot_cursor() as cur:
            cur.execute(sql + " ORDER BY work_date DESC, emp_name, start_min LIMIT 100;", tuple(params))
            rows = [entry_public(r) for r in cur.fetchall()]
    except Exception as e:
        _log_error("check", e)
        return _err("შემოწმება ვერ მოხერხდა", 500)
    return jsonify({"rows": rows})


@bp.route("/api/ot/entries", methods=["POST"])
def create_entries():
    u, err = _require("operator")
    if err:
        return err
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return _err("არასწორი მონაცემები")
    common, errors, warnings = validate_common(data.get("common") if isinstance(data.get("common"), dict) else {})
    emps_in = data.get("employees")
    if not isinstance(emps_in, list) or not emps_in:
        errors.append("აირჩიეთ მინიმუმ ერთი თანამშრომელი")
        emps_in = []
    if len(emps_in) > MAX_BATCH:
        return _err(f"ერთდროულად მაქსიმუმ {MAX_BATCH} თანამშრომელი")
    emps, keys = [], set()
    for e in emps_in:
        if not isinstance(e, dict):
            errors.append("თანამშრომლის მონაცემები არასწორია"); continue
        emp, ee, ww = validate_employee(e, common.get("brigade"))
        errors += ee; warnings += ww
        if emp["key"] in keys:
            errors.append(f"{emp['name']} სიაში ორჯერ არის")
        keys.add(emp["key"])
        emps.append(emp)
    if errors:
        return _err(errors[0], 400, errors=errors)
    confirm = bool(data.get("confirm"))
    s, e_, d = common["start_min"], common["end_min"], common["work_date"]
    day, night = split_day_night(s, e_)
    try:
        with ot_cursor(commit=True) as cur:
            exact_all, confirm_all = [], []
            for emp in emps:
                exact, overlap, same_work = find_conflicts(cur, emp["key"], d, s, e_, common["ganwesi"], common["work_id"])
                for x in exact:
                    exact_all.append({"employee": emp["name"], "kind": "duplicate", "existing": x})
                for x in overlap:
                    confirm_all.append({"employee": emp["name"], "kind": "overlap", "existing": x})
                for x in same_work:
                    confirm_all.append({"employee": emp["name"], "kind": "same_work", "existing": x})
            if exact_all:
                return _err("ზუსტად ასეთი ჩანაწერი უკვე არსებობს — განმეორებით შენახვა შეუძლებელია.", 409,
                            duplicates=exact_all, conflicts=confirm_all, warnings=warnings, blocking=True)
            if (confirm_all or warnings) and not confirm:
                return _err("შენახვამდე გადაამოწმეთ გაფრთხილებები.", 409, duplicates=[], conflicts=confirm_all,
                            warnings=warnings, blocking=False)
            now, created = _now(), []
            for emp in emps:
                uid = uuid.uuid4().hex
                cur.execute(f"""INSERT INTO {T_ENTRIES} (uid, work_date, emp_key, emp_name, emp_pid, emp_position, brigade,
                                work_id, ganwesi, object_name, address, start_min, end_min, day_min, night_min, comment, note,
                                source, source_ref, autofilled, created_by, created_ts)
                                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s);""",
                            (uid, d, emp["key"], emp["name"], emp["pid"], emp["position"], common["brigade"],
                             common["work_id"], common["ganwesi"], common["object_name"], common["address"],
                             s, e_, day, night, common["comment"], common["note"], common["source"],
                             common["source_ref"], common["autofilled"], u["username"], now))
                created.append(uid)
                audit(cur, "create", uid, json.dumps({
                    "employee": emp["name"], "date": d.isoformat(), "ganwesi": common["ganwesi"], "work_id": common["work_id"],
                    "time": f"{fmt_hhmm(s)}–{fmt_hhmm(e_)}", "source": common["source"],
                    "autofilled": common["autofilled"], "confirmed": bool(confirm_all or warnings)}, ensure_ascii=False))
    except Exception as ex:
        if "unique" in str(ex).lower():
            return _err("ზუსტად ასეთი ჩანაწერი უკვე არსებობს (ერთდროული შენახვა).", 409, duplicates=[], conflicts=[], blocking=True)
        _log_error("create", ex)
        return _err("შენახვა ვერ მოხერხდა. სცადეთ თავიდან.", 500)
    return jsonify({"success": True, "created": len(created), "uids": created, "day_min": day, "night_min": night})


def _load_entry(cur, entry_id):
    cur.execute(f"SELECT * FROM {T_ENTRIES} WHERE id = %s;", (int(entry_id),))
    return cur.fetchone()


@bp.route("/api/ot/entries/<int:entry_id>", methods=["PUT"])
def update_entry(entry_id):
    u, err = _require("operator")
    if err:
        return err
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return _err("არასწორი მონაცემები")
    role = role_of(u)
    reason = str(data.get("reason") or "").strip()[:300]
    try:
        with ot_cursor(commit=True) as cur:
            old = _load_entry(cur, entry_id)
            if not old or old.get("deleted_ts"):
                return _err("ჩანაწერი ვერ მოიძებნა", 404)
            if role != "manager" and old["created_by"] != u["username"]:
                return _err("სხვის ჩანაწერს მხოლოდ მენეჯერი ასწორებს", 403)
            allowed = EDITABLE_MANAGER if role == "manager" else EDITABLE_OPERATOR
            forbidden = [k for k in data if k not in allowed and k not in ("reason", "confirm")]
            if forbidden:
                return _err(f"ამ ველების შეცვლის უფლება არ გაქვთ: {', '.join(forbidden)}", 403)
            cur_vals = entry_public(old)
            merged = {
                "work_date": data.get("work_date", cur_vals["work_date"]), "start": data.get("start", cur_vals["start"]),
                "end": data.get("end", cur_vals["end"]), "work_id": data.get("work_id", old["work_id"]),
                "ganwesi": data.get("ganwesi", old["ganwesi"]), "brigade": data.get("brigade", old["brigade"]),
                "object_name": data.get("object_name", old["object_name"]), "address": data.get("address", old["address"]),
                "comment": data.get("comment", old["comment"]), "note": data.get("note", old["note"]),
                "source": old["source"], "source_ref": old["source_ref"], "autofilled": [],
            }
            common, errors, warnings = validate_common(merged)
            emp, ee, ww = validate_employee({"name": data.get("emp_name", old["emp_name"]), "pid": data.get("emp_pid", old["emp_pid"]),
                                             "position": data.get("emp_position", old["emp_position"])}, common.get("brigade"))
            errors += ee
            if errors:
                return _err(errors[0], 400, errors=errors)
            s, e_, d = common["start_min"], common["end_min"], common["work_date"]
            exact, overlap, same_work = find_conflicts(cur, emp["key"], d, s, e_, common["ganwesi"], common["work_id"], exclude_id=entry_id)
            if exact:
                return _err("ზუსტად ასეთი ჩანაწერი უკვე არსებობს.", 409, duplicates=[{"employee": emp["name"], "kind": "duplicate", "existing": x} for x in exact],
                            conflicts=[], warnings=[], blocking=True)
            confirms = [{"employee": emp["name"], "kind": "overlap", "existing": x} for x in overlap]
            if (confirms or warnings) and not data.get("confirm"):
                return _err("შენახვამდე გადაამოწმეთ გაფრთხილებები.", 409, duplicates=[], conflicts=confirms, warnings=warnings, blocking=False)
            day, night = split_day_night(s, e_)
            new_vals = {"work_date": d, "emp_key": emp["key"], "emp_name": emp["name"], "emp_pid": emp["pid"],
                        "emp_position": emp["position"], "brigade": common["brigade"], "work_id": common["work_id"],
                        "ganwesi": common["ganwesi"], "object_name": common["object_name"], "address": common["address"],
                        "start_min": s, "end_min": e_, "day_min": day, "night_min": night,
                        "comment": common["comment"], "note": common["note"]}
            changes = {}
            for k, v in new_vals.items():
                ov = _d(old[k]) if k == "work_date" else old[k]
                if str(ov) != str(v):
                    changes[k] = [str(ov) if k != "emp_pid" else "•••", str(v) if k != "emp_pid" else "•••"]
            if not changes:
                return jsonify({"success": True, "changed": 0})
            sets = ", ".join(f"{k} = %s" for k in new_vals) + ", updated_by = %s, updated_ts = %s"
            cur.execute(f"UPDATE {T_ENTRIES} SET {sets} WHERE id = %s;",
                        tuple(new_vals.values()) + (u["username"], _now(), entry_id))
            audit(cur, "update", old["uid"], json.dumps(changes, ensure_ascii=False), reason)
    except Exception as ex:
        if "unique" in str(ex).lower():
            return _err("ზუსტად ასეთი ჩანაწერი უკვე არსებობს.", 409, duplicates=[], conflicts=[], blocking=True)
        _log_error("update", ex)
        return _err("შენახვა ვერ მოხერხდა. სცადეთ თავიდან.", 500)
    return jsonify({"success": True, "changed": len(changes)})


@bp.route("/api/ot/entries/<int:entry_id>", methods=["DELETE"])
def delete_entry(entry_id):
    u, err = _require("manager")
    if err:
        return err
    reason = str((request.get_json(silent=True) or {}).get("reason") or "").strip()[:300]
    try:
        with ot_cursor(commit=True) as cur:
            old = _load_entry(cur, entry_id)
            if not old or old.get("deleted_ts"):
                return _err("ჩანაწერი ვერ მოიძებნა", 404)
            cur.execute(f"UPDATE {T_ENTRIES} SET deleted_ts = %s, deleted_by = %s WHERE id = %s;", (_now(), u["username"], entry_id))
            audit(cur, "delete", old["uid"], f"{old['emp_name']} · {_d(old['work_date'])} · {old['ganwesi']} · "
                                              f"{fmt_hhmm(old['start_min'])}–{fmt_hhmm(old['end_min'])}", reason)
    except Exception as ex:
        _log_error("delete", ex)
        return _err("წაშლა ვერ მოხერხდა", 500)
    return jsonify({"success": True})


@bp.route("/api/ot/entries/<int:entry_id>/restore", methods=["POST"])
def restore_entry(entry_id):
    u, err = _require("manager")
    if err:
        return err
    try:
        with ot_cursor(commit=True) as cur:
            old = _load_entry(cur, entry_id)
            if not old or not old.get("deleted_ts"):
                return _err("წაშლილი ჩანაწერი ვერ მოიძებნა", 404)
            exact, _, _ = find_conflicts(cur, old["emp_key"], _d(old["work_date"]), old["start_min"], old["end_min"], "", "", exclude_id=entry_id)
            if exact:
                return _err("აღდგენა შეუძლებელია — იგივე ჩანაწერი უკვე არსებობს", 409)
            cur.execute(f"UPDATE {T_ENTRIES} SET deleted_ts = NULL, deleted_by = '' WHERE id = %s;", (entry_id,))
            audit(cur, "restore", old["uid"], f"{old['emp_name']} · {_d(old['work_date'])}")
    except Exception as ex:
        _log_error("restore", ex)
        return _err("აღდგენა ვერ მოხერხდა", 500)
    return jsonify({"success": True})


@bp.route("/api/ot/audit")
def audit_list():
    u, err = _require("operator")
    if err:
        return err
    where, params = [], []
    if role_of(u) != "manager":                       # ოპერატორი ხედავს მხოლოდ საკუთარ ოპერაციებს
        where.append("username = %s"); params.append(u["username"])
    if request.args.get("entry"):
        where.append("entry_uid = %s"); params.append(str(request.args["entry"])[:40])
    sql_where = (" WHERE " + " AND ".join(where)) if where else ""
    try:
        limit = max(1, min(int(request.args.get("limit", 100)), 300))
        offset = max(0, int(request.args.get("offset", 0)))
    except ValueError:
        limit, offset = 100, 0
    try:
        with ot_cursor() as cur:
            cur.execute(f"SELECT COUNT(*) AS n FROM {T_AUDIT}{sql_where};", tuple(params))
            total = int(cur.fetchone()["n"])
            cur.execute(f"SELECT * FROM {T_AUDIT}{sql_where} ORDER BY ts DESC, id DESC LIMIT %s OFFSET %s;",
                        tuple(params) + (limit, offset))
            rows = [{"ts": _ts(r["ts"]), "username": r["username"], "ip": r["ip"], "action": r["action"],
                     "entry_uid": r["entry_uid"], "details": r["details"], "reason": r["reason"]} for r in cur.fetchall()]
    except Exception as e:
        _log_error("audit", e)
        return _err("ისტორიის ჩატვირთვა ვერ მოხერხდა", 500)
    return jsonify({"rows": rows, "total": total})


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
OT_COLUMNS = [
    # (სათაური, სიგანე, ფორმატი)
    ("გვარი, სახელი", 24, "@"), ("პირადი", 26, "@"), ("თარიღი", 11, "dd.mm.yyyy"), ("ბრიგადის #", 9, "General"),
    ("საფუძველი", 11, "General"), ("ობიექტის \nდასახელება", 28, "@"), ("მისამართი", 36, "@"), ("განწესის #", 11, "@"),
    ("დაწყება", 8, "h:mm"), ("დასრულება", 9, "h:mm"),
    ("ბონუს საათი\n(დღე)", 10, "[h]:mm"), ("ბონუს საათი\n(ღამე)", 10, "[h]:mm"), ("შენიშვნა", 12, "@"),
    ("კომენტარი", 18, "@"), ("დამატებითი შენიშვნა", 26, "@"),
]


def build_excel(rows, period_label=""):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    font = "Sylfaen"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "overtime სამონტაჟო"
    head_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    calc_fill = PatternFill(start_color="3B3F8F", end_color="3B3F8F", fill_type="solid")
    night_fill = PatternFill(start_color="EEF0FA", end_color="EEF0FA", fill_type="solid")
    grid = Side(style="thin", color="D6DCE4")
    border = Border(left=grid, right=grid, top=grid, bottom=grid)
    for i, (title, width, _) in enumerate(OT_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = Font(name=font, bold=True, color="FFFFFF", size=9)
        c.fill = calc_fill if i in (11, 12) else head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 40

    body = Font(name=font, size=9)
    for r_i, e in enumerate(rows, start=2):
        d = _d(e["work_date"])
        pid_cell = f"{e['emp_position']} {e['emp_pid']}".strip()
        brig = int(e["brigade"]) if str(e["brigade"]).isdigit() else (e["brigade"] or None)
        wid = int(e["work_id"]) if str(e["work_id"]).isdigit() and len(str(e["work_id"])) < 16 else (e["work_id"] or None)
        vals = [e["emp_name"], pid_cell or None, d, brig, wid, e["object_name"] or None, e["address"] or None,
                e["ganwesi"] or None, time_cls(e["start_min"] // 60, e["start_min"] % 60),
                time_cls(e["end_min"] // 60, e["end_min"] % 60),
                e["day_min"] / 1440, e["night_min"] / 1440, WEEKDAYS[d.weekday()] if d else None,
                e["comment"] or None, e["note"] or None]
        for c_i, v in enumerate(vals, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=v)
            if isinstance(v, str) and cell.data_type == "f":     # ტექსტი „=“-ით ფორმულად არ უნდა შესრულდეს
                cell.data_type = "s"
            cell.number_format = OT_COLUMNS[c_i - 1][2]
            cell.font = body
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if c_i in (1, 2, 6, 7, 14, 15) else "center",
                                       vertical="center", wrap_text=c_i in (6, 7))
            if e["night_min"] and c_i == 12:
                cell.fill = night_fill
    last = len(rows) + 1
    total_row = last + 2
    ws.cell(row=total_row, column=1, value="ჯამი (გაფილტრული)").font = Font(name=font, bold=True, size=10, color="1F3864")
    for col in (11, 12):
        L = get_column_letter(col)
        c = ws.cell(row=total_row, column=col, value=f"=SUBTOTAL(9,{L}2:{L}{last})")
        c.number_format = OT_COLUMNS[col - 1][2]
        c.font = Font(name=font, bold=True, size=10, color="1F3864")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OT_COLUMNS))}{last}"

    # შეჯამება თანამშრომლების მიხედვით (გამოთვლილი სერვერზე — ყველა პროგრამაში სწორად ჩანს)
    sm = wb.create_sheet("შეჯამება")
    heads = [("გვარი, სახელი", 26), ("პირადი", 26), ("ბრიგადა", 12), ("დღეები", 8), ("ჩანაწერები", 10),
             ("დღის საათი", 11), ("ღამის საათი", 11), ("სულ საათი", 11)]
    sm.cell(row=1, column=1, value=f"Overtime — შეჯამება{(' · ' + period_label) if period_label else ''}").font = Font(name=font, bold=True, size=12, color="1F3864")
    for i, (t, w) in enumerate(heads, start=1):
        c = sm.cell(row=3, column=i, value=t)
        c.font = Font(name=font, bold=True, color="FFFFFF", size=9)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sm.column_dimensions[get_column_letter(i)].width = w
    people = {}
    for e in rows:
        p = people.setdefault(e["emp_key"], {"name": e["emp_name"], "pid": f"{e['emp_position']} {e['emp_pid']}".strip(),
                                            "brig": set(), "days": set(), "n": 0, "d": 0, "nt": 0})
        p["brig"].add(str(e["brigade"] or "—")); p["days"].add(_d(e["work_date"])); p["n"] += 1
        p["d"] += e["day_min"]; p["nt"] += e["night_min"]
    r = 4
    for p in sorted(people.values(), key=lambda x: x["name"]):
        vals = [p["name"], p["pid"], ", ".join(sorted(p["brig"])), len(p["days"]), p["n"],
                p["d"] / 1440, p["nt"] / 1440, (p["d"] + p["nt"]) / 1440]
        for c_i, v in enumerate(vals, start=1):
            c = sm.cell(row=r, column=c_i, value=v)
            c.font = body; c.border = border
            c.number_format = "[h]:mm" if c_i >= 6 else "@" if c_i <= 3 else "0"
        r += 1
    sm.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@bp.route("/api/ot/export")
def export_excel():
    u, err = _require("manager")
    if err:
        return err
    sql_where, params = _filters()
    try:
        with ot_cursor() as cur:
            cur.execute(f"SELECT * FROM {T_ENTRIES}{sql_where} ORDER BY work_date, brigade, ganwesi, start_min, emp_name LIMIT 50000;",
                        tuple(params))
            rows = cur.fetchall()
        a = request.args
        label = f"{a.get('from') or '…'} – {a.get('to') or '…'}" if a.get("from") or a.get("to") else ""
        buf = build_excel(rows, label)
        with ot_cursor(commit=True) as cur:
            audit(cur, "export", "", f"{len(rows)} ჩანაწერი" + (f" · {label}" if label else ""))
    except Exception as e:
        _log_error("export", e)
        return _err("Excel-ის გენერირება ვერ მოხერხდა", 500)
    name = f"overtime_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(buf, download_name=name, as_attachment=True,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------------------
# სარეზერვო ასლი, აღდგენა, იმპორტი (მენეჯერი)
# ---------------------------------------------------------------------------
BACKUP_FIELDS = ["uid", "work_date", "emp_key", "emp_name", "emp_pid", "emp_position", "brigade", "work_id", "ganwesi",
                 "object_name", "address", "start_min", "end_min", "day_min", "night_min", "comment", "note", "source",
                 "source_ref", "autofilled", "created_by", "created_ts", "updated_by", "updated_ts", "deleted_ts", "deleted_by"]


@bp.route("/api/ot/backup")
def backup():
    u, err = _require("manager")
    if err:
        return err
    try:
        with ot_cursor() as cur:
            cur.execute(f"SELECT * FROM {T_ENTRIES} ORDER BY id;")
            rows = cur.fetchall()
            out = []
            for r in rows:
                o = {}
                for k in BACKUP_FIELDS:
                    v = r.get(k)
                    o[k] = _d(v).isoformat() if k == "work_date" else (_ts(v) if k.endswith("_ts") else v)
                out.append(o)
        with ot_cursor(commit=True) as cur:
            audit(cur, "backup", "", f"{len(out)} ჩანაწერი")
    except Exception as e:
        _log_error("backup", e)
        return _err("სარეზერვო ასლის შექმნა ვერ მოხერხდა", 500)
    payload = {"module": "overtime", "version": 1, "created": _now().isoformat(), "count": len(out), "entries": out}
    buf = io.BytesIO(json.dumps(payload, ensure_ascii=False).encode("utf-8"))
    return send_file(buf, download_name=f"overtime_backup_{datetime.now().strftime('%Y%m%d_%H%M')}.json",
                     as_attachment=True, mimetype="application/json")


def _insert_rows(cur, items, username, source):
    """არსებულს არ გადაწერს: იგივე uid ან ზუსტი დუბლიკატი → გამოტოვდება."""
    inserted = skipped = 0
    for it in items:
        if it.get("uid"):
            cur.execute(f"SELECT 1 AS x FROM {T_ENTRIES} WHERE uid = %s;", (it["uid"],))
            if cur.fetchone():
                skipped += 1; continue
        if not it.get("deleted_ts"):
            cur.execute(f"""SELECT 1 AS x FROM {T_ENTRIES} WHERE emp_key = %s AND work_date = %s AND start_min = %s
                            AND end_min = %s AND deleted_ts IS NULL;""", (it["emp_key"], it["work_date"], it["start_min"], it["end_min"]))
            if cur.fetchone():
                skipped += 1; continue
        cur.execute(f"""INSERT INTO {T_ENTRIES} (uid, work_date, emp_key, emp_name, emp_pid, emp_position, brigade, work_id,
                        ganwesi, object_name, address, start_min, end_min, day_min, night_min, comment, note, source,
                        source_ref, autofilled, created_by, created_ts, updated_by, updated_ts, deleted_ts, deleted_by)
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s);""",
                    (it.get("uid") or uuid.uuid4().hex, it["work_date"], it["emp_key"], it["emp_name"], it.get("emp_pid", ""),
                     it.get("emp_position", ""), it.get("brigade", ""), it.get("work_id", ""), it.get("ganwesi", ""),
                     it.get("object_name", ""), it.get("address", ""), it["start_min"], it["end_min"], it["day_min"],
                     it["night_min"], it.get("comment", ""), it.get("note", ""), source, it.get("source_ref", ""),
                     it.get("autofilled", ""), it.get("created_by") or username, it.get("created_ts") or _now(),
                     it.get("updated_by", ""), it.get("updated_ts") or None, it.get("deleted_ts") or None, it.get("deleted_by", "")))
        inserted += 1
    return inserted, skipped


def _clean_item(raw, strict_uid=False):
    """იმპორტის/აღდგენის ერთი სტრიქონი → შესანახი ფორმა, ან (None, მიზეზი)."""
    if not isinstance(raw, dict):
        return None, "არასწორი სტრიქონი"
    name = norm_name(raw.get("emp_name"))
    d = parse_date(raw.get("work_date"))
    s = raw.get("start_min") if isinstance(raw.get("start_min"), int) else parse_hhmm(raw.get("start"))
    e = raw.get("end_min") if isinstance(raw.get("end_min"), int) else parse_hhmm(raw.get("end"))
    if not name or not d or s is None or e is None or not (0 <= s < 1440 and 0 <= e < 1440) or s == e:
        return None, "აკლია სახელი, თარიღი ან დრო"
    pid = re.sub(r"\D", "", str(raw.get("emp_pid") or ""))[:20]
    day, night = split_day_night(s, e)
    item = {"emp_name": name[:120], "emp_pid": pid, "emp_position": str(raw.get("emp_position") or "")[:80],
            "emp_key": emp_key(name, pid), "work_date": d, "start_min": s, "end_min": e, "day_min": day, "night_min": night,
            "brigade": str(raw.get("brigade") or "").strip()[:40], "work_id": norm_work_id(raw.get("work_id")),
            "ganwesi": norm_ganwesi(raw.get("ganwesi")), "object_name": str(raw.get("object_name") or "")[:300],
            "address": str(raw.get("address") or "")[:400], "comment": str(raw.get("comment") or "")[:300],
            "note": str(raw.get("note") or "")[:500], "source_ref": str(raw.get("source_ref") or "")[:100]}
    if strict_uid:
        uid = str(raw.get("uid") or "")
        if not re.fullmatch(r"[0-9a-f]{32}", uid):
            return None, "არასწორი uid"
        item.update(uid=uid, created_by=str(raw.get("created_by") or "")[:100], created_ts=raw.get("created_ts") or None,
                    updated_by=str(raw.get("updated_by") or "")[:100], updated_ts=raw.get("updated_ts") or None,
                    deleted_ts=raw.get("deleted_ts") or None, deleted_by=str(raw.get("deleted_by") or "")[:100],
                    autofilled=str(raw.get("autofilled") or "")[:300])
    return item, ""


@bp.route("/api/ot/restore", methods=["POST"])
def restore_backup():
    """სარეზერვო JSON-ის აღდგენა: მხოლოდ ამატებს დაკარგულ ჩანაწერებს, არსებულს არ ცვლის და არ შლის."""
    u, err = _require("manager")
    if err:
        return err
    data = request.get_json(silent=True)
    entries = (data or {}).get("entries") if isinstance(data, dict) else None
    if not isinstance(entries, list) or not entries or len(entries) > MAX_IMPORT_BATCH:
        return _err(f"გამოგზავნეთ 1–{MAX_IMPORT_BATCH} ჩანაწერი ერთ ნაწილად")
    items, bad = [], 0
    for raw in entries:
        it, why = _clean_item(raw, strict_uid=True)
        if it:
            items.append(it)
        else:
            bad += 1
    try:
        with ot_cursor(commit=True) as cur:
            ins, skip = _insert_rows(cur, items, u["username"], "restore")
            audit(cur, "restore_backup", "", f"დაემატა {ins}, გამოტოვდა {skip}, არასწორი {bad}")
    except Exception as e:
        _log_error("restore_backup", e)
        return _err("აღდგენა ვერ მოხერხდა", 500)
    return jsonify({"success": True, "inserted": ins, "skipped": skip, "invalid": bad})


@bp.route("/api/ot/import", methods=["POST"])
def import_rows():
    """overtime.xlsx-ის ძველი სტრიქონები (ბრაუზერში წაკითხული, ნაწილ-ნაწილ). დუბლიკატები გამოტოვდება."""
    u, err = _require("manager")
    if err:
        return err
    data = request.get_json(silent=True)
    rows = (data or {}).get("rows") if isinstance(data, dict) else None
    if not isinstance(rows, list) or not rows or len(rows) > MAX_IMPORT_BATCH:
        return _err(f"გამოგზავნეთ 1–{MAX_IMPORT_BATCH} სტრიქონი ერთ ნაწილად")
    items, bad, seen, dup_in_batch = [], [], set(), 0
    for i, raw in enumerate(rows):
        it, why = _clean_item(raw)
        if not it:
            bad.append({"row": raw.get("row") if isinstance(raw, dict) else i, "reason": why}); continue
        k = (it["emp_key"], it["work_date"], it["start_min"], it["end_min"])
        if k in seen:
            dup_in_batch += 1
            continue
        seen.add(k); items.append(it)
    try:
        with ot_cursor(commit=True) as cur:
            ins, skip = _insert_rows(cur, items, u["username"], "import")
            skip += dup_in_batch
            audit(cur, "import", "", f"დაემატა {ins}, დუბლიკატი {skip}, გამოტოვებული {len(bad)}")
    except Exception as e:
        _log_error("import", e)
        return _err("იმპორტი ვერ მოხერხდა", 500)
    return jsonify({"success": True, "inserted": ins, "duplicates": skip, "invalid": bad[:50], "invalid_count": len(bad)})


# ---------------------------------------------------------------------------
def init_app(app, deps):
    """ძირითადი აპლიკაციიდან ერთხელ: მოდულის რეგისტრაცია.
    ბაზასთან კავშირი აქ არ ხდება — სქემა პირველ Overtime-მოთხოვნაზე იქმნება,
    ასე რომ Overtime-ის ბაზის პრობლემა ძირითადი აპლიკაციის გაშვებას არ აფერხებს."""
    global _deps
    _deps = deps
    app.register_blueprint(bp)
